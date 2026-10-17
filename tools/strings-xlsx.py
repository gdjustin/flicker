"""Turn dist/strings.json into the owner's text spreadsheet.
Usage: python3 tools/strings-xlsx.py dist/strings.json dist/flicker-text-v1r.xlsx"""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

rows = json.load(open(sys.argv[1]))
out = sys.argv[2]
ORDER = ['Modals and key moments', 'Journal', 'Prologue', 'Cards and upgrades', 'Settlers and resources', 'Hovers and buttons',
         'Map and places', 'Rebirth and perks', 'Time away', 'Text in the screen code', 'Other']
ABOUT = {
    'Modals and key moments': 'The pop-up moments: the village opening, first settler, the map reveal, the amplifier, the thirty-place rebirth prompt.',
    'Journal': 'Every line the game writes to the journal (the log on the right), with when it appears.',
    'Prologue': 'Becoming a self: the parts, the upgrades, the titles and their messages.',
    'Cards and upgrades': 'Ideas, research cards, buildings and lines, projects and their stages, in the village and the map phase.',
    'Settlers and resources': 'Jobs (name, short tag, hover), resource names, gathering buttons, looking around.',
    'Hovers and buttons': 'Hover text, button labels and small headings across the village and the drawer.',
    'Map and places': 'The map phase: the sending strip, place cards, relics, land and building-group names.',
    'Rebirth and perks': 'The rebirth (now called scatter in places), the between-runs screen, perk medallions and comforts.',
    'Time away': 'The "while you were away" report and its lines.',
    'Text in the screen code': 'Words written straight into the screen code. {…} marks a number or name the game fills in. Edit the sentence as a whole; the code line is there for context.',
    'Other': 'Anything that fit no other tab.'}

F = 'Arial'
head_font = Font(name=F, bold=True, color='FFFFFF', size=10)
head_fill = PatternFill('solid', start_color='3A3027')
body = Font(name=F, size=10)
muted = Font(name=F, size=9, color='75685A')
edit_fill = PatternFill('solid', start_color='FFF2B3')
flag_font = Font(name=F, size=9, color='A3341F', bold=True)
thin = Side(style='thin', color='D6CBB6')
wrap = Alignment(wrap_text=True, vertical='top')

wb = Workbook()
ws = wb.active; ws.title = 'Read me'
ws.column_dimensions['A'].width = 34; ws.column_dimensions['B'].width = 70; ws.column_dimensions['C'].width = 16; ws.column_dimensions['D'].width = 16
lines = [
    ('Flicker (working title): every word in the game, v1r, 2026-10-17', None),
    ('', None),
    ('How to edit', 'Type your new wording in the yellow "Your edit" column. Leave it empty to keep the current text. Use "Notes" for anything else (cut it, merge it, ask a question).'),
    ('Keep the {…} and {name} parts', 'Words in curly brackets are filled in by the game (a number, a place name). Keep them where they belong, or say in Notes if one should go.'),
    ('Flags (red)', 'tile, grid, steps: break the fiction rule. scatter → rebirth?: the prestige word. "(cut)": the system is cut, so the text never shows now; no edit needed unless it comes back.'),
    ('ID and Source', 'Tell Claude where each line lives, so your edits can be folded back. Don\'t change them.'),
    ('Example', 'Current text "Pick a tile on the map." → Your edit "Pick a place on the map." (Flags: tile)'),
    ('', None),
    ('Tab', 'What it holds'),
]
r = 1
for a, b in lines:
    ws.cell(r, 1, a).font = Font(name=F, size=14 if r == 1 else 10, bold=(r == 1 or b is not None and a in ('Tab',)))
    if b: ws.cell(r, 2, b).font = Font(name=F, size=10, bold=(a == 'Tab')); ws.cell(r, 2).alignment = wrap
    if a and b and a != 'Tab': ws.cell(r, 1).font = Font(name=F, size=10, bold=True)
    r += 1
ws.cell(r - 1, 3, 'Rows').font = Font(name=F, size=10, bold=True)
ws.cell(r - 1, 4, 'Edited so far').font = Font(name=F, size=10, bold=True)

by = {}
for x in rows: by.setdefault(x['sheet'] if x['sheet'] in ORDER else 'Other', []).append(x)
for name in ORDER:
    items = by.get(name)
    if not items: continue
    items.sort(key=lambda x: (x.get('where', ''), x.get('file', ''), x.get('line', 0)))
    sh = wb.create_sheet(name[:31])
    cols = ['ID', 'Where', 'What it is']
    if name == 'Journal': cols += ['When it appears']
    cols += ['Current text', 'Your edit', 'Notes', 'Flags', 'Source']
    if name == 'Text in the screen code': cols += ['Code line']
    widths = {'ID': 14, 'Where': 22, 'What it is': 26, 'When it appears': 34, 'Current text': 60, 'Your edit': 60, 'Notes': 28, 'Flags': 16, 'Source': 22, 'Code line': 60}
    for i, c in enumerate(cols, 1):
        cell = sh.cell(1, i, c); cell.font = head_font; cell.fill = head_fill; cell.alignment = Alignment(vertical='center', wrap_text=True)
        sh.column_dimensions[get_column_letter(i)].width = widths[c]
    for j, x in enumerate(items, 2):
        vals = {'ID': x['id'], 'Where': x.get('where', ''), 'What it is': x.get('what', ''), 'When it appears': x.get('when', ''), 'Current text': x['text'], 'Your edit': None,
                'Notes': None, 'Flags': x.get('flags', ''), 'Source': (x.get('file', '') + (':' + str(x['line']) if x.get('line') else '')), 'Code line': x.get('code', '')}
        for i, c in enumerate(cols, 1):
            cell = sh.cell(j, i, vals[c]); cell.alignment = wrap; cell.border = Border(bottom=thin)
            cell.font = muted if c in ('ID', 'Source', 'Code line') else flag_font if c == 'Flags' else body
            if c == 'Your edit': cell.fill = edit_fill
    sh.freeze_panes = 'E2' if name != 'Journal' else 'F2'
    sh.auto_filter.ref = 'A1:' + get_column_letter(len(cols)) + str(len(items) + 1)
    edit_col = get_column_letter(cols.index('Your edit') + 1)
    ws.cell(r, 1, name).font = Font(name=F, size=10, bold=True)
    ws.cell(r, 2, ABOUT[name]).font = body; ws.cell(r, 2).alignment = wrap
    ws.cell(r, 3, len(items)).font = body
    ws.cell(r, 4, "=COUNTA('%s'!%s2:%s%d)" % (sh.title, edit_col, edit_col, len(items) + 1)).font = body
    r += 1
ws.cell(r, 1, 'All tabs').font = Font(name=F, size=10, bold=True)
ws.cell(r, 3, '=SUM(C%d:C%d)' % (r - len([n for n in ORDER if n in by]), r - 1)).font = Font(name=F, size=10, bold=True)
ws.cell(r, 4, '=SUM(D%d:D%d)' % (r - len([n for n in ORDER if n in by]), r - 1)).font = Font(name=F, size=10, bold=True)
wb.save(out)
print('wrote', out, len(rows), 'rows')
